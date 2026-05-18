#!/usr/bin/env python3
"""Reads Protocolo.xlsx and generates index.html for GitHub Pages deployment."""

import json
import openpyxl
from pathlib import Path

EXCEL_FILE = "Protocolo.xlsx"
OUTPUT_FILE = "index.html"

# Protocolo sheet — column indices (0-based)
P_NOMBRE = 2        # C: product name
P_INGREDIENTE = 3   # D: active ingredient
P_CONCENTRACION = 4 # E: concentration
P_TIPO_PRES = 9     # J: presentation type (COMPRIMIDO, JARABE, etc.)
P_UNIDAD = 10       # K: unit (CAJA DE 30, FRASCO, etc.)
P_ES_PROTOCOLO = 12 # M: 1 = in protocol

# Novedades sheet — column indices (0-based, data starts col B = index 1)
N_NOMBRE = 3        # D: product name
N_INGREDIENTE = 4   # E: active ingredient
N_CONCENTRACION = 5 # F: concentration
N_TIPO_PRES = 10    # K: presentation type
N_UNIDAD = 11       # L: unit


def read_protocolo(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Protocolo"]
    productos = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[P_ES_PROTOCOLO] != 1:
            continue
        nombre = str(row[P_NOMBRE] or "").strip()
        if not nombre:
            continue
        ingrediente = str(row[P_INGREDIENTE] or "").strip()
        concentracion = str(row[P_CONCENTRACION] or "").strip()
        tipo = str(row[P_TIPO_PRES] or "").strip()
        unidad = str(row[P_UNIDAD] or "").strip()
        presentacion = f"{tipo} · {unidad}" if tipo and unidad else tipo or unidad
        productos.append({
            "nombre": nombre,
            "ingrediente": ingrediente,
            "concentracion": concentracion,
            "presentacion": presentacion,
        })
    return sorted(productos, key=lambda x: x["ingrediente"])


def read_novedades(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Novedades"]
    novedades = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row[N_NOMBRE]:
            continue
        nombre = str(row[N_NOMBRE]).strip()
        ingrediente = str(row[N_INGREDIENTE] or "").strip().title()
        concentracion = str(row[N_CONCENTRACION] or "").strip()
        tipo = str(row[N_TIPO_PRES] or "").strip().title()
        unidad = str(row[N_UNIDAD] or "").strip().title()
        pills = [p for p in [concentracion, tipo, unidad] if p]
        novedades.append({"nombre": nombre, "ingrediente": ingrediente, "pills": pills})
    return novedades


def build_novedad_cards(novedades):
    cards = []
    for n in novedades:
        pills_html = "".join(
            f'<span class="novedad-pill">{p}</span>' for p in n["pills"]
        )
        cards.append(
            f'      <div class="novedad-card">\n'
            f'        <div class="novedad-producto">{n["nombre"]}</div>\n'
            f'        <div class="novedad-ingrediente">{n["ingrediente"]}</div>\n'
            f'        <div class="novedad-pills">{pills_html}</div>\n'
            f'      </div>'
        )
    return "\n".join(cards)


def generate(excel_path=EXCEL_FILE, output_path=OUTPUT_FILE):
    productos = read_protocolo(excel_path)
    novedades = read_novedades(excel_path)
    novedad_cards = build_novedad_cards(novedades)
    productos_json = json.dumps(productos, ensure_ascii=False)

    html = HTML_TEMPLATE.replace("{{NOVEDAD_CARDS}}", novedad_cards)
    html = html.replace("{{PRODUCTOS_JSON}}", productos_json)

    Path(output_path).write_text(html, encoding="utf-8")
    print(f"Generated {output_path}: {len(productos)} products, {len(novedades)} novedades.")


HTML_TEMPLATE = """\
<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Protocolo de Ingredientes Activos</title>
<link href="https://fonts.googleapis.com/css2?family=DM+Sans:wght@300;400;500;600&family=DM+Mono:wght@400;500&display=swap" rel="stylesheet">
<style>
  *, *::before, *::after { box-sizing: border-box; margin: 0; padding: 0; }
  :root {
    --teal: #0F6E56; --teal-light: #E1F5EE; --teal-mid: #1D9E75;
    --amber: #BA7517; --amber-light: #FAEEDA;
    --coral: #993C1D; --coral-light: #FAECE7;
    --blue: #185FA5; --blue-light: #E6F1FB;
    --gray: #5F5E5A; --gray-light: #F1EFE8; --gray-mid: #D3D1C7;
    --text: #2C2C2A; --text-muted: #888780;
    --bg: #fafaf8; --white: #ffffff;
    --radius: 10px; --radius-sm: 6px;
  }
  body { font-family: 'DM Sans', sans-serif; background: var(--bg); color: var(--text); min-height: 100vh; }
  .header {
    background: var(--white); border-bottom: 1px solid var(--gray-mid);
    padding: 1.25rem 2rem; display: flex; align-items: center;
    justify-content: space-between; position: sticky; top: 0; z-index: 100;
  }
  .header-brand { display: flex; align-items: center; gap: 10px; }
  .header-icon {
    width: 36px; height: 36px; background: var(--teal); border-radius: 8px;
    display: flex; align-items: center; justify-content: center;
  }
  .header-icon svg { width: 20px; height: 20px; fill: white; }
  .header-title { font-size: 1.05rem; font-weight: 600; color: var(--text); }
  .header-sub { font-size: 0.72rem; color: var(--text-muted); font-weight: 400; margin-top: 1px; }
  .header-count { background: var(--teal-light); color: var(--teal); font-size: 0.78rem; font-weight: 500; padding: 4px 12px; border-radius: 20px; }
  .main { max-width: 1100px; margin: 0 auto; padding: 2rem 1.5rem 3rem; }
  .section-label { font-size: 0.68rem; font-weight: 600; letter-spacing: 0.1em; text-transform: uppercase; color: var(--text-muted); margin-bottom: 0.75rem; }
  .novedades-banner {
    background: linear-gradient(135deg, #0F6E56 0%, #1D9E75 100%);
    border-radius: var(--radius); padding: 1.5rem; margin-bottom: 2rem;
    position: relative; overflow: hidden;
  }
  .novedades-banner::before {
    content: ''; position: absolute; top: -30px; right: -30px;
    width: 140px; height: 140px; border-radius: 50%; background: rgba(255,255,255,0.06);
  }
  .novedades-banner::after {
    content: ''; position: absolute; bottom: -50px; right: 80px;
    width: 200px; height: 200px; border-radius: 50%; background: rgba(255,255,255,0.04);
  }
  .novedades-header { display: flex; align-items: center; gap: 10px; margin-bottom: 1.25rem; }
  .badge-new { background: rgba(255,255,255,0.2); color: white; font-size: 0.68rem; font-weight: 600; letter-spacing: 0.08em; text-transform: uppercase; padding: 3px 10px; border-radius: 20px; border: 1px solid rgba(255,255,255,0.3); }
  .novedades-title { color: white; font-size: 1.05rem; font-weight: 600; }
  .novedades-grid { display: grid; grid-template-columns: 1fr 1fr; gap: 12px; position: relative; z-index: 1; }
  @media (max-width: 580px) { .novedades-grid { grid-template-columns: 1fr; } }
  .novedad-card { background: rgba(255,255,255,0.12); border: 1px solid rgba(255,255,255,0.2); border-radius: var(--radius-sm); padding: 1rem 1.25rem; backdrop-filter: blur(4px); }
  .novedad-producto { font-size: 0.72rem; color: rgba(255,255,255,0.65); margin-bottom: 4px; }
  .novedad-ingrediente { color: white; font-size: 1rem; font-weight: 600; margin-bottom: 6px; }
  .novedad-pills { display: flex; gap: 6px; flex-wrap: wrap; }
  .novedad-pill { background: rgba(255,255,255,0.15); color: rgba(255,255,255,0.9); font-size: 0.72rem; font-weight: 500; padding: 2px 9px; border-radius: 12px; font-family: 'DM Mono', monospace; }
  .sugerencias-box { background: var(--white); border: 1px solid var(--gray-mid); border-radius: var(--radius); padding: 1.25rem 1.5rem; margin-bottom: 2rem; display: flex; align-items: center; gap: 16px; }
  .sugerencias-icon { flex-shrink: 0; width: 42px; height: 42px; background: var(--amber-light); border-radius: 10px; display: flex; align-items: center; justify-content: center; }
  .sugerencias-icon svg { width: 22px; height: 22px; fill: var(--amber); }
  .sugerencias-text h3 { font-size: 0.92rem; font-weight: 600; color: var(--text); margin-bottom: 3px; }
  .sugerencias-text p { font-size: 0.82rem; color: var(--text-muted); }
  .sugerencias-btn { margin-left: auto; flex-shrink: 0; background: var(--amber); color: white; font-size: 0.82rem; font-weight: 600; padding: 8px 18px; border-radius: 8px; text-decoration: none; white-space: nowrap; transition: background 0.15s; }
  .sugerencias-btn:hover { background: var(--coral); }
  .toolbar { display: flex; gap: 10px; margin-bottom: 1.25rem; flex-wrap: wrap; }
  .search-wrap { flex: 1; min-width: 200px; position: relative; }
  .search-wrap svg { position: absolute; left: 12px; top: 50%; transform: translateY(-50%); width: 16px; height: 16px; stroke: var(--text-muted); fill: none; stroke-width: 2; stroke-linecap: round; }
  .search-input { width: 100%; padding: 8px 12px 8px 36px; border: 1px solid var(--gray-mid); border-radius: var(--radius-sm); font-size: 0.88rem; font-family: 'DM Sans', sans-serif; background: var(--white); color: var(--text); outline: none; transition: border-color 0.15s; }
  .search-input:focus { border-color: var(--teal-mid); }
  .filter-select { padding: 8px 12px; border: 1px solid var(--gray-mid); border-radius: var(--radius-sm); font-size: 0.85rem; font-family: 'DM Sans', sans-serif; background: var(--white); color: var(--text); outline: none; cursor: pointer; }
  .results-info { font-size: 0.8rem; color: var(--text-muted); padding: 4px 0; align-self: center; }
  .table-wrap { background: var(--white); border: 1px solid var(--gray-mid); border-radius: var(--radius); overflow: hidden; }
  table { width: 100%; border-collapse: collapse; }
  thead { background: var(--gray-light); }
  th { font-size: 0.72rem; font-weight: 600; letter-spacing: 0.06em; text-transform: uppercase; color: var(--gray); padding: 11px 16px; text-align: left; border-bottom: 1px solid var(--gray-mid); }
  td { padding: 11px 16px; font-size: 0.875rem; border-bottom: 1px solid #f0efe8; vertical-align: middle; }
  tr:last-child td { border-bottom: none; }
  tr:hover td { background: #f7f6f1; }
  .td-nombre { font-size: 0.8rem; color: var(--text-muted); }
  .td-ing { font-weight: 500; color: var(--text); }
  .td-conc { font-family: 'DM Mono', monospace; font-size: 0.8rem; color: var(--blue); background: var(--blue-light); padding: 3px 8px; border-radius: 5px; display: inline-block; }
  .td-pres { font-size: 0.8rem; color: var(--gray); }
  .no-results { text-align: center; padding: 3rem; color: var(--text-muted); font-size: 0.9rem; }
  .pagination { display: flex; align-items: center; justify-content: center; gap: 6px; padding: 1.25rem 0 0.5rem; }
  .page-btn { width: 34px; height: 34px; border: 1px solid var(--gray-mid); border-radius: 6px; background: var(--white); font-size: 0.85rem; cursor: pointer; color: var(--text); display: flex; align-items: center; justify-content: center; transition: all 0.12s; }
  .page-btn:hover { border-color: var(--teal-mid); color: var(--teal); }
  .page-btn.active { background: var(--teal); color: white; border-color: var(--teal); }
  .page-btn:disabled { opacity: 0.35; cursor: default; }
</style>
</head>
<body>

<header class="header">
  <div class="header-brand">
    <div class="header-icon">
      <svg viewBox="0 0 24 24"><path d="M9 3H5a2 2 0 00-2 2v4m6-6h10a2 2 0 012 2v4M9 3v18m0 0h10a2 2 0 002-2V9M9 21H5a2 2 0 01-2-2V9m0 0h18"/></svg>
    </div>
    <div>
      <div class="header-title">Protocolo Farmac&#233;utico</div>
      <div class="header-sub">Ingredientes activos en protocolo</div>
    </div>
  </div>
  <div class="header-count" id="total-count"></div>
</header>

<main class="main">

  <div class="section-label">&#10022; Novedades</div>
  <div class="novedades-banner">
    <div class="novedades-header">
      <span class="badge-new">Nuevos ingresos</span>
      <span class="novedades-title">Reci&#233;n incorporados al protocolo</span>
    </div>
    <div class="novedades-grid">
{{NOVEDAD_CARDS}}
    </div>
  </div>

  <div class="sugerencias-box">
    <div class="sugerencias-icon">
      <svg viewBox="0 0 24 24"><path d="M8 10h.01M12 10h.01M16 10h.01M9 16H5a2 2 0 01-2-2V6a2 2 0 012-2h14a2 2 0 012 2v8a2 2 0 01-2 2h-5l-5 5v-5z"/></svg>
    </div>
    <div class="sugerencias-text">
      <h3>Sugerencias de ingredientes activos</h3>
      <p>&#191;Tienes alguna sugerencia para ampliar el protocolo? Comparte tu recomendaci&#243;n con el equipo.</p>
    </div>
    <a class="sugerencias-btn" href="https://forms.office.com/r/FR4jxLmiZc?origin=lprLink" target="_blank">
      Enviar sugerencia &#8594;
    </a>
  </div>

  <div class="section-label">Cat&#225;logo completo</div>
  <div class="toolbar">
    <div class="search-wrap">
      <svg viewBox="0 0 24 24"><circle cx="11" cy="11" r="8"/><line x1="21" y1="21" x2="16.65" y2="16.65"/></svg>
      <input class="search-input" type="text" id="searchInput" placeholder="Buscar ingrediente activo, producto..." oninput="applyFilters()">
    </div>
    <select class="filter-select" id="presFilter" onchange="applyFilters()">
      <option value="">Todas las presentaciones</option>
    </select>
    <span class="results-info" id="resultsInfo"></span>
  </div>

  <div class="table-wrap">
    <table>
      <thead>
        <tr>
          <th style="width:35%">Ingrediente activo</th>
          <th style="width:18%">Concentraci&#243;n</th>
          <th style="width:22%">Presentaci&#243;n</th>
          <th style="width:25%">Producto</th>
        </tr>
      </thead>
      <tbody id="tableBody"></tbody>
    </table>
    <div id="noResults" class="no-results" style="display:none">No se encontraron resultados para tu b&#250;squeda.</div>
  </div>

  <div class="pagination" id="pagination"></div>

</main>

<script>
const PRODUCTOS = {{PRODUCTOS_JSON}};
const PER_PAGE = 25;
let filtered = [...PRODUCTOS];
let page = 1;
function cap(s) { return s ? s.charAt(0).toUpperCase() + s.slice(1).toLowerCase() : ''; }
function titleCase(s) { return s.split(' ').map(w => cap(w)).join(' '); }

function buildPresOptions() {
  const types = [...new Set(PRODUCTOS.map(p => p.presentacion.split(' · ')[0]).filter(Boolean))].sort();
  const sel = document.getElementById('presFilter');
  types.forEach(t => {
    const o = document.createElement('option');
    o.value = t; o.textContent = titleCase(t);
    sel.appendChild(o);
  });
}

function applyFilters() {
  const q = document.getElementById('searchInput').value.toLowerCase().trim();
  const pf = document.getElementById('presFilter').value;
  filtered = PRODUCTOS.filter(p => {
    const matchQ = !q || p.ingrediente.toLowerCase().includes(q) || p.nombre.toLowerCase().includes(q) || p.concentracion.toLowerCase().includes(q);
    const matchP = !pf || p.presentacion.startsWith(pf);
    return matchQ && matchP;
  });
  page = 1;
  render();
}

function render() {
  const start = (page - 1) * PER_PAGE;
  const slice = filtered.slice(start, start + PER_PAGE);
  const tbody = document.getElementById('tableBody');
  const noRes = document.getElementById('noResults');
  document.getElementById('resultsInfo').textContent = filtered.length + ' resultado' + (filtered.length !== 1 ? 's' : '');
  if (filtered.length === 0) {
    tbody.innerHTML = '';
    noRes.style.display = 'block';
  } else {
    noRes.style.display = 'none';
    tbody.innerHTML = slice.map(p => `
      <tr>
        <td><div class="td-ing">${titleCase(p.ingrediente)}</div></td>
        <td>${p.concentracion ? '<span class="td-conc">' + p.concentracion + '</span>' : '<span style="color:var(--text-muted);font-size:0.8rem">—</span>'}</td>
        <td class="td-pres">${titleCase(p.presentacion.replace('·', '·'))}</td>
        <td class="td-nombre">${titleCase(p.nombre)}</td>
      </tr>`).join('');
  }
  renderPagination();
}

function renderPagination() {
  const total = Math.ceil(filtered.length / PER_PAGE);
  const pag = document.getElementById('pagination');
  if (total <= 1) { pag.innerHTML = ''; return; }
  let html = `<button class="page-btn" onclick="goPage(${page-1})" ${page===1?'disabled':''}>&#8592;</button>`;
  const range = [];
  for (let i = 1; i <= total; i++) {
    if (i === 1 || i === total || (i >= page - 2 && i <= page + 2)) range.push(i);
    else if (range[range.length-1] !== '…') range.push('…');
  }
  range.forEach(r => {
    if (r === '…') html += `<span style="padding:0 4px;color:var(--text-muted)">…</span>`;
    else html += `<button class="page-btn ${r===page?'active':''}" onclick="goPage(${r})">${r}</button>`;
  });
  html += `<button class="page-btn" onclick="goPage(${page+1})" ${page===total?'disabled':''}>&#8594;</button>`;
  pag.innerHTML = html;
}

function goPage(n) {
  const total = Math.ceil(filtered.length / PER_PAGE);
  if (n < 1 || n > total) return;
  page = n;
  render();
  document.querySelector('.table-wrap').scrollIntoView({behavior:'smooth', block:'start'});
}

document.getElementById('total-count').textContent = PRODUCTOS.length + ' productos';
buildPresOptions();
render();
</script>
</body>
</html>
"""

if __name__ == "__main__":
    generate()
